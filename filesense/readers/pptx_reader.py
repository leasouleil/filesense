from pptx import Presentation

from filesense.readers.base import BaseReader


class PptxReader(BaseReader):

    def read(self, filepath: str) -> dict:
        presentation = Presentation(filepath)

        text = ""

        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text += shape.text + "\n"

        return {
            "text": text,
            "metadata": {},
        }